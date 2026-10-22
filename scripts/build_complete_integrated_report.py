# -*- coding: utf-8 -*-
"""
build_complete_integrated_report.py
Builds the complete, final, professional dissertation for MCTE:
1. Section 1 (Cover Page & DS Comments, elements 0 to 102):
   - Preserves official MCTE Academy Cover Page & DS Comments pages from
     PROJECT_REPORT_FILE(TNM&RSR).docx (title banners, shapes, textboxes with
     OC Tanmay Mugale TES-4312 & OC Rishabh Rathore TES-4316, DS Guide Lt Col Rajat Gaur,
     DS Co-Guide Maj Vaibhav Kukreti, and DS Comments).
2. Section 2 (Front Matter, starting directly on Page 4):
   - NO duplicate section break (prevents blank Page 4).
   - Starts immediately with '2. ACKNOWLEDGEMENTS' on Page 4.
   - Front matter headings ('2. ACKNOWLEDGEMENTS', '3. ABSTRACT', '4. TABLE OF CONTENTS',
     '5. LIST OF FIGURES', '6. LIST OF TABLES', '7. LIST OF ABBREVIATIONS AND ACRONYMS')
     are strictly centered and underlined with Navy bold styling.
   - Table of Contents is cleanly formatted with bulleted and indented entries:
     '•\t1. Title Page & DS Comments', '•\t2. Acknowledgements', '•\t3. Abstract',
     bold chapter titles, and 0.40" indented subsections without raw asterisks or dynamic field codes.
3. Equations:
   - All 39 display equations ($$...$$) are converted into native Microsoft Word OMML (<m:oMath>)
     via latex2mathml and MML2OMML.XSL, centered and beautifully rendered.
   - All inline math expressions ($...$) are sanitized with unicode math symbols (λ, θ, π, μ, σ,
     ², ³, ×, ≈, →, ₁, ₂, ₃, ₄), completely removing raw LaTeX markup.
4. Header & Numbering:
   - Top-centered continuous Arabic page numbers (1, 2, 3...) throughout all pages.
   - Empty footers and no top-right header text.
5. Saves exclusively to:
   - E:\WIFI_INTERFEROMETRIC_NEURAL_SENSING\PROJECT_REPORT_FILE(TNM&RSR).docx
"""
import os
import re
import sys
from pathlib import Path

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

import latex2mathml.converter
from lxml import etree

REPO_ROOT = Path(r"e:\WIFI_INTERFEROMETRIC_NEURAL_SENSING\wimotion_2.0")
PARENT_DIR = Path(r"e:\WIFI_INTERFEROMETRIC_NEURAL_SENSING")
MD_PATH = REPO_ROOT / "docs" / "MASTER_DISSERTATION_FINAL.md"
TNM_BASE_PATH = PARENT_DIR / "PROJECT_REPORT_FILE(TNM&RSR).docx"

GENERATED_DIAGRAMS_DIR = REPO_ROOT / "docs" / "figures" / "generated_diagrams"
EXTRACTED_PHOTOS_DIR = REPO_ROOT / "docs" / "figures" / "extracted_photos"
ASSETS_DIR = REPO_ROOT / "docs" / "assets"

DIAGRAM_DEFINITIONS = [
    {
        "matcher": lambda b: "PROJECT ECHO" in b and "PROJECT SHADOW" in b and "WIMOTION 1.0" in b,
        "filename": "diagram_1_project_evolution.png",
        "caption_label": "Figure 1.3: ",
        "caption_desc": "Architectural Evolution Across WiMotion Development Phases (Phases 1 to 4)"
    },
    {
        "matcher": lambda b: "[INCIDENT EM WAVE]" in b,
        "filename": "diagram_propagation_mechanisms.png",
        "caption_label": "Figure 2.1: ",
        "caption_desc": "Four Fundamental Electromagnetic Propagation Mechanisms in Indoor Architectural Environments"
    },
    {
        "matcher": lambda b: "RSSI (Coarse Scalar Measurement)" in b,
        "filename": "diagram_2_rssi_vs_csi.png",
        "caption_label": "Figure 2.2: ",
        "caption_desc": "Structural Comparison Between Scalar RSSI and Orthogonal 64-Subcarrier Channel State Information (CSI)"
    },
    {
        "matcher": lambda b: "1st Fresnel Zone Boundary" in b or "d1 + d2 = D + 1ambda/2" in b or "d1 + d2 = D + lambda/2" in b,
        "filename": "diagram_3_fresnel_zone_geometry.png",
        "caption_label": "Figure 2.3: ",
        "caption_desc": "First Fresnel Zone Ellipsoid Geometry and Operating Boundary Parameters"
    },
    {
        "matcher": lambda b: "[SINGLE-LINK AMBIGUITY]" in b,
        "filename": "diagram_4_single_vs_multilink.png",
        "caption_label": "Figure 2.4: ",
        "caption_desc": "Comparison Between Single-Link Longitudinal Ambiguity and 8-Link Spatial Mesh Intersection"
    },
    {
        "matcher": lambda b: "ESP32-WROOM-32 SoC" in b and "Xtensa LX6 Core 0" in b,
        "filename": "diagram_esp32_hardware_arch.png",
        "caption_label": "Figure 3.1: ",
        "caption_desc": "Hardware Architecture and Baseband CSI Extraction of the ESP32-WROOM-32 Node"
    },
    {
        "matcher": lambda b: "[TX2: Zone 2 (NW)]" in b and "[TX1: Zone 1 (NE)]" in b,
        "filename": "diagram_5_star_topology_mesh.png",
        "caption_label": "Figure 3.2: ",
        "caption_desc": "Star-Topology Multi-Link Sensing Mesh: Distributed Corner Transmitters and Central Dual Receivers"
    },
    {
        "matcher": lambda b: "[TX2: Node NW]" in b and "Northern Wall" in b,
        "filename": "diagram_room_coordinates_layout.png",
        "caption_label": "Figure 3.3: ",
        "caption_desc": "Spatial Coordinate Grid and Quadrant Partitioning of the ROOM_LAB_A Testbed"
    },
    {
        "matcher": lambda b: "[PHYSICAL LAYER]" in b and "COM8 @ 921.6k" in b,
        "filename": "diagram_7_software_pipeline_architecture.png",
        "caption_label": "Figure 4.1: ",
        "caption_desc": "Five-Tier Multithreaded Serial Ingestion and CSI Signal Processing Pipeline Architecture"
    },
    {
        "matcher": lambda b: "[STAGE 1: RAW BUFFER]" in b,
        "filename": "diagram_8_canonical_feature_pipeline.png",
        "caption_label": "Figure 5.1: ",
        "caption_desc": "Four-Stage Canonical Spatial Feature Extraction and Fusion Pipeline"
    },
    {
        "matcher": lambda b: "[FASTAPI APPLICATION SERVER]" in b,
        "filename": "diagram_9_client_server_visualization.png",
        "caption_label": "Figure 7.1: ",
        "caption_desc": "Decoupled Client-Server Software Architecture and Real-Time Telemetry Pipeline"
    },
    {
        "matcher": lambda b: "VALID CSI RATE" in b and "ACTIVE LINKS" in b and "RECEIVER HEALTH" in b,
        "filename": "diagram_telemetry_hud_cards.png",
        "caption_label": "Figure 7.2: ",
        "caption_desc": "Real-Time Diagnostic Telemetry Metric Cards and Hardware Health Indicators"
    },
    {
        "matcher": lambda b: "SAFE STANDBY" in b,
        "filename": "diagram_10_fsm_state_machine.png",
        "caption_label": "Figure 7.3: ",
        "caption_desc": "Finite State Machine (FSM) Transition Model of the WiMotion Replay Lab and Simulation Environment"
    },
    {
        "matcher": lambda b: "[TIER 1: AUTOMATED UNIT TESTS]" in b,
        "filename": "diagram_11_four_tier_verification.png",
        "caption_label": "Figure 8.1: ",
        "caption_desc": "Four-Tier Software Verification, Continuous Integration, and Cryptographic Provenance Hierarchy"
    },
    {
        "matcher": lambda b: "[EXPERIMENTAL SCALE vs. OPERATIONAL APPLICATION]" in b,
        "filename": "diagram_12_prototype_vs_operational_scope.png",
        "caption_label": "Figure 9.1: ",
        "caption_desc": "System Scope Boundary: Laboratory Experimental Prototype vs Operational Border Perimeter Deployment"
    },
    {
        "matcher": lambda b: "WIMOTION STRUCTURED ENGINEERING ROADMAP" in b,
        "filename": "diagram_13_engineering_roadmap.png",
        "caption_label": "Figure 10.1: ",
        "caption_desc": "Four-Stage Engineering Upgrade Roadmap for Next-Generation Wi-Fi Sensing"
    }
]

FIGURE_IMAGE_MAP = {
    "FIGURE 1.1": EXTRACTED_PHOTOS_DIR / "image1.png",
    "FIGURE 1.2": EXTRACTED_PHOTOS_DIR / "image2.png",
    "FIGURE 3.1": EXTRACTED_PHOTOS_DIR / "image3.jpg",
    "FIGURE 3.2": EXTRACTED_PHOTOS_DIR / "image4.jpg",
    "FIGURE 3.3": EXTRACTED_PHOTOS_DIR / "image5.jpg",
    "FIGURE 3.4": EXTRACTED_PHOTOS_DIR / "image6.jpg",
    "FIGURE 3.5": EXTRACTED_PHOTOS_DIR / "image7.jpg",
    "FIGURE 3.6": EXTRACTED_PHOTOS_DIR / "image8.jpg",
    "FIGURE 6.1": ASSETS_DIR / "fig_6_1_confusion_matrix.png",
    "FIGURE 6.2": ASSETS_DIR / "fig_6_2_model_comparison.png",
    "FIGURE 6.3": ASSETS_DIR / "fig_6_3_per_zone_recall.png",
    "FIGURE 6.4": ASSETS_DIR / "fig_6_4_false_alarm_reduction.png",
    "FIGURE 7.1": ASSETS_DIR / "fig_7_1_live_observatory.png",
    "FIGURE 7.2": ASSETS_DIR / "fig_7_2_replay_lab.png",
    "FIGURE 7.3": EXTRACTED_PHOTOS_DIR / "image11.png",
    "FIGURE 8.1": EXTRACTED_PHOTOS_DIR / "image10.png",
    "FIGURE H.1": EXTRACTED_PHOTOS_DIR / "image10.png",
    "FIGURE H.2": EXTRACTED_PHOTOS_DIR / "image9.jpg",
    "FIGURE H.3": EXTRACTED_PHOTOS_DIR / "image8.jpg",
}

COLOR_NAVY = RGBColor(27, 54, 93)       # #1B365D - Deep Navy
COLOR_SLATE = RGBColor(51, 65, 85)      # #334155 - Slate Grey
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B - Muted Grey
COLOR_BODY = RGBColor(30, 41, 59)       # #1E293B - Off-Black

# Load XSLT stylesheet for MathML -> OMML
XSLT_PATH = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
try:
    XSLT_DOC = etree.parse(XSLT_PATH)
    MML_TRANSFORM = etree.XSLT(XSLT_DOC)
    MATH_AVAILABLE = True
except Exception as e:
    print(f"[!] Warning: MML2OMML.XSL could not be loaded ({e}). Math will fallback to unicode.")
    MATH_AVAILABLE = False


def configure_styles(doc):
    styles = doc.styles

    try:
        norm = styles['Normal']
        norm.font.name = 'Times New Roman'
        norm.font.size = Pt(12)
        norm.font.color.rgb = COLOR_BODY
        norm.paragraph_format.line_spacing = 1.35
        norm.paragraph_format.space_after = Pt(6)
        norm.paragraph_format.space_before = Pt(0)
    except Exception as e:
        print("Note on Normal style:", e)

    h_configs = [
        ('Heading 1', 16, True, False, COLOR_NAVY, 18, 8),
        ('Heading 2', 14, True, False, COLOR_SLATE, 14, 6),
        ('Heading 3', 12, True, True, COLOR_SLATE, 10, 4),
    ]

    for name, sz, b, it, col, sb, sa in h_configs:
        try:
            h = styles[name]
        except KeyError:
            h = styles.add_style(name, docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
        h.font.name = 'Times New Roman'
        h.font.size = Pt(sz)
        h.font.bold = b
        h.font.italic = it
        h.font.color.rgb = col
        h.paragraph_format.space_before = Pt(sb)
        h.paragraph_format.space_after = Pt(sa)
        h.paragraph_format.keep_with_next = True


def apply_page_margins(section):
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)


def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def clean_latex_string(expr):
    """Fallback text cleaner for math expressions."""
    replacements = [
        (r'\\text\{([^}]+)\}', r'\1'),
        (r'\\mathbf\{([^}]+)\}', r'\1'),
        (r'\\mathit\{([^}]+)\}', r'\1'),
        (r'\\mathrm\{([^}]+)\}', r'\1'),
        (r'\\mathbb\{I\}', 'I'),
        (r'\\approx', '≈'),
        (r'\\times', '×'),
        (r'\\cdot', '·'),
        (r'\\pm', '±'),
        (r'\\to', '→'),
        (r'\\rightarrow', '→'),
        (r'\\Rightarrow', '⇒'),
        (r'\\lambda', 'λ'),
        (r'\\theta', 'θ'),
        (r'\\pi', 'π'),
        (r'\\mu', 'μ'),
        (r'\\sigma', 'σ'),
        (r'\\Delta', 'Δ'),
        (r'\\Phi', 'Φ'),
        (r'\\tau', 'τ'),
        (r'\\phi', 'φ'),
        (r'\\le(q)?', '≤'),
        (r'\\ge(q)?', '≥'),
        (r'\\neq', '≠'),
        (r'\\in', '∈'),
        (r'\\sum', '∑'),
        (r'\\sqrt\{([^}]+)\}', r'√(\1)'),
        (r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)'),
        (r'\\left\(', '('),
        (r'\\right\)', ')'),
        (r'\\left\[', '['),
        (r'\\right\]', ']'),
        (r'\\quad', '   '),
        (r'\\,', ' '),
        (r'\\;', ' '),
        (r'\\!', ''),
        (r'_1', '₁'),
        (r'_2', '₂'),
        (r'_3', '₃'),
        (r'_4', '₄'),
        (r'_k', 'ₖ'),
        (r'_t', 'ₜ'),
        (r'_r', 'ᵣ'),
        (r'_m', 'ₘ'),
        (r'_0', '₀'),
        (r'_i', 'ᵢ'),
        (r'\^2', '²'),
        (r'\^3', '³'),
        (r'\^8', '⁸'),
        (r'\^9', '⁹'),
        (r'\{', ''),
        (r'\}', ''),
        (r'\\', ''),
    ]
    res = expr
    for pat, repl in replacements:
        res = re.sub(pat, repl, res)
    return res.strip()


def clean_inline_math(text):
    """Sanitizes $formula$ expressions inside paragraph text into clean Unicode math."""
    def _repl(m):
        raw = m.group(1).strip()
        return clean_latex_string(raw)
    return re.sub(r'\$([^\$]+)\$', _repl, text)


def render_math_equation(doc, latex_str):
    """Renders a display equation ($$...$$) as native Word OMML (<m:oMath>)."""
    eq = latex_str.strip()
    if eq.startswith("$$"):
        eq = eq[2:]
    if eq.endswith("$$"):
        eq = eq[:-2]
    eq = eq.strip()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    rendered_omml = False
    if MATH_AVAILABLE:
        try:
            mml = latex2mathml.converter.convert(eq)
            mml_tree = etree.fromstring(mml.encode('utf-8'))
            omml_tree = MML_TRANSFORM(mml_tree)
            omml_xml = etree.tostring(omml_tree.getroot())
            omml_elem = parse_xml(omml_xml)
            p._p.append(omml_elem)
            rendered_omml = True
        except Exception as e:
            rendered_omml = False

    if not rendered_omml:
        clean_text = clean_latex_string(eq)
        run = p.add_run(clean_text)
        run.font.name = "Cambria Math"
        run.font.size = Pt(11.5)
        run.italic = True
        run.font.color.rgb = COLOR_NAVY


def add_formatted_text(paragraph, text, font_size=12, default_bold=False, default_italic=False, default_color=COLOR_BODY):
    """Processes markdown tokens (**bold**, *italic*, `code`) and inline math."""
    text = clean_inline_math(text)
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for tok in tokens:
        if not tok:
            continue
        is_bold = default_bold
        is_italic = default_italic
        is_code = False
        tok_clean = tok

        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            is_bold = True
            tok_clean = tok[2:-2]
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            is_italic = True
            tok_clean = tok[1:-1]
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            is_code = True
            tok_clean = tok[1:-1]

        run = paragraph.add_run(tok_clean)
        run.font.size = Pt(font_size)
        run.bold = is_bold
        run.italic = is_italic
        if is_code:
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            run.font.name = "Times New Roman"
            run.font.color.rgb = default_color


def render_code_block(doc, code_lines):
    code_text = "\n".join(code_lines)

    # Check if this matches one of the 16 publication technical diagrams
    matched_diag = None
    for d in DIAGRAM_DEFINITIONS:
        if d["matcher"](code_text):
            matched_diag = d
            break

    if matched_diag:
        img_path = GENERATED_DIAGRAMS_DIR / matched_diag["filename"]
        if img_path.exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            r_img = p_img.add_run()
            r_img.add_picture(str(img_path), width=Inches(5.8))

            cap_p = doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(3)
            cap_p.paragraph_format.space_after = Pt(8)

            r_label = cap_p.add_run(matched_diag["caption_label"])
            r_label.font.name = "Times New Roman"
            r_label.font.size = Pt(10)
            r_label.bold = True
            r_label.font.color.rgb = COLOR_NAVY

            r_desc = cap_p.add_run(matched_diag["caption_desc"])
            r_desc.font.name = "Times New Roman"
            r_desc.font.size = Pt(10)
            r_desc.italic = True
            r_desc.font.color.rgb = COLOR_BODY
            return

    # Regular code block or terminal output log
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F5F7FA")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="1B365D"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'</w:tblBorders>'
    )
    tbl._tbl.tblPr.append(borders)

    cp = cell.paragraphs[0]
    cp.paragraph_format.line_spacing = 1.05
    cp.paragraph_format.space_after = Pt(0)
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = cp.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 40, 50)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


def render_image_placeholder(doc, line):
    caption_text = re.sub(r'[\*\[\]]', '', line).strip()
    if caption_text.startswith("- ") or caption_text.startswith("* "):
        caption_text = caption_text[2:].strip()

    m_fig = re.search(r'(Figure\s+[0-9A-Za-z\.]+)', caption_text, re.IGNORECASE)
    fig_tag = m_fig.group(1).upper() if m_fig else "FIGURE"

    # Check if a real image file exists for this figure
    img_file = FIGURE_IMAGE_MAP.get(fig_tag)
    if img_file and Path(img_file).exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        r_img = p_img.add_run()
        r_img.add_picture(str(img_file), width=Inches(5.8))

        cap_p = doc.add_paragraph()
        cap_p.paragraph_format.space_before = Pt(3)
        cap_p.paragraph_format.space_after = Pt(8)
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        parts = caption_text.split(":", 1)
        if len(parts) == 2:
            clean_lbl = clean_inline_math(parts[0].strip()) + ": "
            r_label = cap_p.add_run(clean_lbl)
            r_label.font.name = "Times New Roman"
            r_label.font.size = Pt(10)
            r_label.bold = True
            r_label.font.color.rgb = COLOR_NAVY

            clean_desc = clean_inline_math(parts[1].strip())
            r_desc = cap_p.add_run(clean_desc)
            r_desc.font.name = "Times New Roman"
            r_desc.font.size = Pt(10)
            r_desc.italic = True
            r_desc.font.color.rgb = COLOR_BODY
        else:
            clean_desc = clean_inline_math(caption_text)
            r_desc = cap_p.add_run(clean_desc)
            r_desc.font.name = "Times New Roman"
            r_desc.font.size = Pt(10)
            r_desc.bold = True
            r_desc.font.color.rgb = COLOR_NAVY
        return True

    # Otherwise render placeholder table
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=240, bottom=240, left=320, right=320)

    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'</w:tblBorders>'
    )
    tbl._tbl.tblPr.append(borders)

    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.line_spacing = 1.2
    cp.paragraph_format.space_after = Pt(2)

    r1 = cp.add_run(f"📷  [ INSERT {fig_tag} HERE ]\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.bold = True
    r1.font.color.rgb = COLOR_NAVY

    r2 = cp.add_run("Suggested Format: High-Resolution Landscape Image (Width: ~5.5 inches)\n(Word Action: Click inside this box ➔ Insert ➔ Pictures ➔ This Device)")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.0)
    r2.italic = True
    r2.font.color.rgb = COLOR_MUTED

    cap_p = doc.add_paragraph()
    cap_p.paragraph_format.space_before = Pt(4)
    cap_p.paragraph_format.space_after = Pt(8)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    parts = caption_text.split(":", 1)
    if len(parts) == 2:
        clean_lbl = clean_inline_math(parts[0].strip()) + ": "
        r_label = cap_p.add_run(clean_lbl)
        r_label.font.name = "Times New Roman"
        r_label.font.size = Pt(10)
        r_label.bold = True
        r_label.font.color.rgb = COLOR_NAVY

        clean_desc = clean_inline_math(parts[1].strip())
        r_desc = cap_p.add_run(clean_desc)
        r_desc.font.name = "Times New Roman"
        r_desc.font.size = Pt(10)
        r_desc.italic = True
        r_desc.font.color.rgb = COLOR_BODY
    else:
        clean_desc = clean_inline_math(caption_text)
        r_desc = cap_p.add_run(clean_desc)
        r_desc.font.name = "Times New Roman"
        r_desc.font.size = Pt(10)
        r_desc.bold = True
        r_desc.font.color.rgb = COLOR_NAVY

    return True


def render_table(doc, table_buffer):
    if not table_buffer or len(table_buffer) < 2:
        return

    cleaned_rows = []
    for r in table_buffer:
        r_str = r.strip()
        if r_str.startswith("|"):
            r_str = r_str[1:]
        if r_str.endswith("|"):
            r_str = r_str[:-1]
        cols = [c.strip() for c in r_str.split("|")]
        cleaned_rows.append(cols)

    if not cleaned_rows:
        return

    header_cols = cleaned_rows[0]
    data_rows = []
    for r in cleaned_rows[1:]:
        if all(re.match(r'^:?-+:?$', c) for c in r if c):
            continue
        data_rows.append(r)

    num_cols = len(header_cols)
    if num_cols == 0:
        return

    num_rows = 1 + len(data_rows)
    tbl = doc.add_table(rows=num_rows, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="1B365D"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="1B365D"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl._tbl.tblPr.append(borders)

    for c_idx, col_text in enumerate(header_cols):
        cell = tbl.cell(0, c_idx)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(col_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_cols in enumerate(data_rows):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx in range(num_cols):
            cell = tbl.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            txt = row_cols[c_idx] if c_idx < len(row_cols) else ""
            add_formatted_text(p, txt, font_size=9.5, default_color=COLOR_BODY)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


def build_integrated_dissertation():
    print(f"[*] Reading master markdown: {MD_PATH}")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    # Load base template document
    print(f"[*] Loading base document: {TNM_BASE_PATH}")
    doc = Document(str(TNM_BASE_PATH))
    configure_styles(doc)

    # Clean up base document: keep elements 0 to 102 (Cover Page, textboxes & DS comments)
    # and preserve final root sectPr (body[-1])
    body = doc._body._element
    print(f"[*] Initial base elements: {len(body)}")
    while len(body) > 104:
        body.remove(body[103])
    print(f"[+] Preserved official cover & comments. Trimmed elements to: {len(body)}")

    # Clean any MCEME occurrences in preserved base elements
    for el in body[:103]:
        for t in el.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
            if 'MCEME' in t.text or 'Mechanical' in t.text:
                t.text = t.text.replace('Military College of Electronics and Mechanical Engineering (MCEME)', 'Military College of Telecommunication Engineering (MCTE)')
                t.text = t.text.replace('MCEME', 'MCTE')

    # SECTION 1: Cover Page & DS Comments (elements 0 to 102)
    sec_cover = doc.sections[0]
    apply_page_margins(sec_cover)
    sec_cover.different_first_page_header_footer = True
    for p in sec_cover.footer.paragraphs:
        p.text = ""

    lines = content.splitlines()
    total_lines = len(lines)
    i = 0

    in_code_block = False
    code_buffer = []
    table_buffer = []
    in_table = False
    in_front_matter_figure_list = False
    current_section = 2  # Section 1 is Cover (index 0). Front matter is Section 2 (index 1).

    # Skip lines until "## 2. ACKNOWLEDGEMENTS" in markdown because Cover Page & DS comments are already in base doc
    while i < total_lines:
        if "## 2. ACKNOWLEDGEMENTS" in lines[i]:
            break
        i += 1

    # NOTE: Do NOT call doc.add_section() here!
    # Element 102 already contains the sectPr ending Section 0 (Cover & DS comments).
    # Starting paragraphs directly at element 103 places Acknowledgements directly on Page 4,
    # completely eliminating the empty Page 4!

    front_matter_page_breaks = [
        "## 2. ACKNOWLEDGEMENTS",
        "## 3. ABSTRACT",
        "## 4. TABLE OF CONTENTS",
        "## 5. LIST OF FIGURES",
        "## 6. LIST OF TABLES",
        "## 7. LIST OF ABBREVIATIONS"
    ]

    while i < total_lines:
        line = lines[i]
        line_stripped = line.strip()

        # Transition from Front Matter to Chapter 1 (Section 3)
        if line_stripped.startswith("# CHAPTER 1:") and current_section == 2:
            current_section = 3
            sec_chapters = doc.add_section()
            apply_page_margins(sec_chapters)
            sec_chapters.different_first_page_header_footer = False
            sec_chapters.header.is_linked_to_previous = True
            sec_chapters.footer.is_linked_to_previous = True

        # Code Blocks (```)
        if line_stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                render_code_block(doc, code_buffer)
                code_buffer = []
            else:
                in_code_block = True
                code_buffer = []
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        # Tables (| col | col |)
        if "|" in line_stripped and (line_stripped.startswith("|") or line_stripped.endswith("|")):
            in_table = True
            table_buffer.append(line_stripped)
            i += 1
            continue
        else:
            if in_table:
                render_table(doc, table_buffer)
                table_buffer = []
                in_table = False

        # Display Math Equations ($$...$$)
        if line_stripped.startswith("$$") and line_stripped.endswith("$$") and len(line_stripped) > 2:
            render_math_equation(doc, line_stripped)
            i += 1
            continue

        # Horizontal rules / Page Break delimiters
        if line_stripped in ("---", "***", "___") or "page-break-after" in line_stripped:
            if current_section == 2:
                next_line = ""
                for k in range(i + 1, min(i + 6, total_lines)):
                    if lines[k].strip():
                        next_line = lines[k].strip()
                        break

                if any(sec in next_line for sec in front_matter_page_breaks):
                    doc.add_page_break()

            i += 1
            continue

        # Track Front Matter List of Figures
        if "5. LIST OF FIGURES" in line_stripped or "LIST OF FIGURES" in line_stripped:
            in_front_matter_figure_list = True
        elif in_front_matter_figure_list and ("LIST OF TABLES" in line_stripped or line_stripped == "---"):
            in_front_matter_figure_list = False

        # Figure Placeholders & Captions
        if not in_front_matter_figure_list and re.match(r'^[\*\-\s]*(\*\*)?\[?Figure\s+[A-Za-z0-9\.]+', line_stripped, re.IGNORECASE):
            if render_image_placeholder(doc, line_stripped):
                i += 1
                continue

        # Front Matter Headings: Centered and Underlined
        if any(line_stripped.startswith(prefix) for prefix in [
            "## 2. ACKNOWLEDGEMENTS",
            "## 3. ABSTRACT",
            "## 4. TABLE OF CONTENTS",
            "## 5. LIST OF FIGURES",
            "## 6. LIST OF TABLES",
            "## 7. LIST OF ABBREVIATIONS"
        ]):
            h_text = re.sub(r'^##\s*', '', line_stripped).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.bold = True
            run.underline = True
            run.font.color.rgb = COLOR_NAVY

            # If this is Table of Contents, render Tanmay's custom-formatted TOC list
            if "TABLE OF CONTENTS" in h_text:
                i += 1
                # Consume and format all TOC bullet lines until next section break (---)
                while i < total_lines:
                    toc_l = lines[i].strip()
                    if toc_l in ("---", "***", "___") or toc_l.startswith("## 5."):
                        break
                    if not toc_l:
                        i += 1
                        continue

                    # Skip unwanted items
                    if any(term in toc_l for term in ["Keywords", "Appendix I:", "Appendix J:", "Certificate of Approval", "Student Declaration"]):
                        i += 1
                        continue

                    # Check if chapter / part title or subsection based on raw line indentation
                    raw_toc_l = lines[i]
                    is_chapter = False
                    clean_toc_text = ""
                    if re.match(r'^\s{2,}[\-\*]', raw_toc_l) or raw_toc_l.startswith('\t'):
                        # Indented subsection or front matter sub-item
                        is_chapter = False
                        clean_toc_text = re.sub(r'^\s*[\-\*]\s*', '', toc_l).strip()
                    elif toc_l.startswith("- **") or toc_l.startswith("* **"):
                        is_chapter = True
                        clean_toc_text = re.sub(r'^[\-\*]\s*\*\*', '', toc_l)
                        clean_toc_text = re.sub(r'\*\*.*$', '', clean_toc_text).strip()
                    elif toc_l.startswith("- ") or toc_l.startswith("* "):
                        is_chapter = True
                        clean_toc_text = re.sub(r'^[\-\*]\s*', '', toc_l).strip()
                    else:
                        is_chapter = False
                        clean_toc_text = re.sub(r'^\s*[\-\*]\s*', '', toc_l).strip()

                    # Specific rename requested by user: Title Page & DS Comments
                    if "Title Page" in clean_toc_text:
                        clean_toc_text = "1. Title Page & DS Comments"

                    clean_toc_text = clean_inline_math(clean_toc_text)
                    clean_toc_text = clean_toc_text.replace("**", "").replace("*", "")

                    tp = doc.add_paragraph()
                    tp.paragraph_format.line_spacing = 1.15
                    if is_chapter:
                        tp.paragraph_format.left_indent = Inches(0.15)
                        tp.paragraph_format.space_before = Pt(5)
                        tp.paragraph_format.space_after = Pt(2)
                        r = tp.add_run(f"•\t{clean_toc_text}")
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(11.5)
                        r.bold = True
                        r.font.color.rgb = COLOR_NAVY
                    else:
                        tp.paragraph_format.left_indent = Inches(0.40)
                        tp.paragraph_format.space_before = Pt(1)
                        tp.paragraph_format.space_after = Pt(1)
                        r = tp.add_run(f"•\t{clean_toc_text}")
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(11.0)
                        r.bold = False
                        r.font.color.rgb = COLOR_SLATE
                    i += 1
                continue

            i += 1
            continue

        # Skip any Keywords section
        if "KEYWORDS" in line_stripped.upper() and (line_stripped.startswith("##") or line_stripped.startswith("#")):
            while i + 1 < total_lines:
                next_l = lines[i + 1].strip()
                if next_l.startswith("#") or next_l == "---":
                    break
                i += 1
            i += 1
            continue

        # Skip Appendix I or Appendix J
        if any(app in line_stripped.upper() for app in ["APPENDIX I:", "APPENDIX I ", "APPENDIX J:", "APPENDIX J "]) and line_stripped.startswith("##"):
            while i + 1 < total_lines:
                next_l = lines[i + 1].strip()
                if (next_l.startswith("## APPENDIX") or next_l.startswith("# ")) and not any(app in next_l.upper() for app in ["APPENDIX I", "APPENDIX J"]):
                    break
                i += 1
            i += 1
            continue

        # Heading 1: Chapters and Part Titles
        if line_stripped.startswith("# "):
            h_text = line_stripped[2:].strip()
            p = doc.add_paragraph(style='Heading 1')
            if current_section == 3 and not line_stripped.startswith("# CHAPTER 1:"):
                p.paragraph_format.page_break_before = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ("PART" in h_text or "REFERENCES" in h_text) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        # Heading 2: Sections
        if line_stripped.startswith("## "):
            h_text = line_stripped[3:].strip()
            p = doc.add_paragraph(style='Heading 2')
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = COLOR_SLATE
            i += 1
            continue

        # Heading 3: Subsections
        if line_stripped.startswith("### "):
            h_text = line_stripped[4:].strip()
            p = doc.add_paragraph(style='Heading 3')
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.bold = True
            run.font.italic = True
            run.font.color.rgb = COLOR_SLATE
            i += 1
            continue

        # Blockquote / Callout (> [!NOTE] or > Text)
        if line_stripped.startswith(">"):
            b_text = re.sub(r'^>\s*(\[!.*?\])?', '', line_stripped).strip()
            if b_text:
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F0F4F8")
                set_cell_margins(cell, top=100, bottom=100, left=180, right=140)

                borders = parse_xml(
                    f'<w:tblBorders {nsdecls("w")}>'
                    f'<w:left w:val="single" w:sz="18" w:space="0" w:color="1B365D"/>'
                    f'<w:top w:val="none"/>'
                    f'<w:bottom w:val="none"/>'
                    f'<w:right w:val="none"/>'
                    f'</w:tblBorders>'
                )
                tbl._tbl.tblPr.append(borders)

                cp = cell.paragraphs[0]
                cp.paragraph_format.line_spacing = 1.2
                cp.paragraph_format.space_after = Pt(0)
                add_formatted_text(cp, b_text, font_size=11, default_italic=True, default_color=COLOR_SLATE)
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(3)

            i += 1
            continue

        # Bullet lists (- or *)
        if line_stripped.startswith("- ") or line_stripped.startswith("* "):
            item_text = line_stripped[2:].strip()
            # Filter out unwanted items
            if any(term in item_text for term in ["Keywords", "Appendix I:", "Appendix J:", "Certificate of Approval", "Student Declaration"]):
                i += 1
                continue
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, item_text, font_size=11.5)
            i += 1
            continue

        # Numbered lists (1. or 2.)
        m_num = re.match(r'^(\d+)\.\s+(.*)$', line_stripped)
        if m_num:
            num_str = m_num.group(1)
            item_text = m_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(2)
            nrun = p.add_run(f"{num_str}.  ")
            nrun.font.name = "Times New Roman"
            nrun.font.size = Pt(11.5)
            nrun.bold = True
            nrun.font.color.rgb = COLOR_SLATE
            add_formatted_text(p, item_text, font_size=11.5)
            i += 1
            continue

        # Empty lines
        if not line_stripped:
            i += 1
            continue

        # Standard Body Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_text(p, line_stripped, font_size=11.5)
        i += 1

    if in_table and table_buffer:
        render_table(doc, table_buffer)

    # =========================================================================
    # POST-PROCESSING: Strict Header & Page Number Cleanup Across All Sections
    # =========================================================================
    print("[*] Running strict header and page number post-processing...")
    for idx, sec in enumerate(doc.sections):
        sectPr = sec._sectPr
        # Remove any pgNumType so numbering is continuous Arabic (1, 2, 3...) throughout
        for pg in sectPr.findall(qn('w:pgNumType')):
            sectPr.remove(pg)

        if idx == 0:
            sec.different_first_page_header_footer = True
        else:
            sec.different_first_page_header_footer = False
            sec.header.is_linked_to_previous = True
            sec.footer.is_linked_to_previous = True
            # Remove any explicit headerReference/footerReference so Word strictly inherits from Section 0
            for hr in sectPr.findall(qn('w:headerReference')):
                sectPr.remove(hr)
            for fr in sectPr.findall(qn('w:footerReference')):
                sectPr.remove(fr)

        # Clear any text runs in headers and footers
        for p in sec.header.paragraphs:
            for r in p.runs:
                if r.text and any(w in r.text for w in ['WiMotion', 'Dissertation', 'Sensing', 'Final', 'Multi-Link']):
                    r.text = ''
        for p in sec.footer.paragraphs:
            for r in p.runs:
                r.text = ''

    # Clean any residual text in all related parts (header2.xml, header3.xml, footer2.xml, etc.)
    for rId, part in doc.part.related_parts.items():
        if 'header' in part.partname or 'footer' in part.partname:
            if hasattr(part, '_element'):
                for p in part._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                    for t in p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                        if t.text and any(w in t.text for w in ['WiMotion', 'Dissertation', 'Sensing', 'Final', 'Page']):
                            t.text = ''

    # Save exclusively to the final report destination: PROJECT_REPORT_FILE(TNM&RSR).docx
    final_target = TNM_BASE_PATH

    print(f"[*] Saving finalized integrated dissertation to: {final_target}")
    doc.save(str(final_target))
    print(f"[+] Saved final report: {final_target} ({os.path.getsize(final_target):,} bytes)")
    print("[+] All done successfully!")


if __name__ == "__main__":
    build_integrated_dissertation()
